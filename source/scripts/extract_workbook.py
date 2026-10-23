"""Read the supplied XLSX without altering it; export traced, cleaned case records.

Only this XLSX is an input. No PDF, website state, or prior model output is used.
"""
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
import json
import math
import re
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SOURCE = Path(sys.argv[1] if len(sys.argv)>1 else 'CaseQuest2026Dataset.xlsx')
OUT = Path(sys.argv[2]) if len(sys.argv)>2 else Path(__file__).parent
OUT.mkdir(parents=True, exist_ok=True)
STAGES = ['detect_hours', 'dispatch_hours', 'parts_wait_hours', 'repair_hours', 'restart_hours']
KEYS = {'Assets': ['asset_id'], 'Events': ['event_id'], 'Production': ['asset_id', 'month'],
        'Readings': ['reading_id'], 'Solutions': ['option_id'], 'HSE': ['hse_id']}
DATE_FIELDS = {'month', 'due_date', 'closed_date', 'status_as_of'}
NUMERIC_FIELDS = {
    'Assets': ['age_years','modeled_capacity_bbl_hr','vibration_normal_mm_s','temperature_normal_c','replacement_usd','network_uptime_pct','technicians','spare_fill_pct','training_hours_per_person','edge_setup_usd','edge_annual_usd','cleanup_setup_usd','available_project_hours'],
    'Events': STAGES + ['stopped_hours','lost_oil_bbl','recoverable_fraction','labor_hours','labor_rate_usd_hr','parts_usd','contractor_usd','protected_repair_hours'],
    'Production': ['calendar_hours','potential_oil_bbl','produced_oil_bbl','gross_deferred_bbl','recovered_oil_bbl','price_usd_bbl','variable_cost_usd_bbl'],
    'Readings': ['vibration_mm_s','temperature_c'],
    'HSE': ['potential_severity','likelihood_band','action_setup_usd','action_annual_usd','action_person_hours','spilled_oil_bbl','recovered_spill_bbl'],
    'Solutions': ['setup_usd','annual_usd','lead_months','install_person_hours'] + [stage+'_'+band for stage in ['detect','dispatch','parts_wait','repair','restart'] for band in ['low','base','high']],
}
cleaning = []
schema = []
tables = {}

def plain(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v

def numeric(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    match = re.fullmatch(r'\s*(-?\d+(?:\.\d+)?)\s*(?:bbl/hr|F|C)?\s*', str(v).replace(',', ''), re.I)
    if not match:
        raise ValueError('Unrecognized numeric value: '+repr(v))
    number = float(match[1])
    return int(number) if number.is_integer() else number

def normalize(sheet, field, v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
    if field.endswith('_id'):
        v = str(v).strip().upper()
        if field == 'asset_id' and re.fullmatch(r'A?\d+', v):
            v = 'A'+v.removeprefix('A').zfill(3)
    if field in NUMERIC_FIELDS.get(sheet, []):
        if field == 'temperature_c' and isinstance(v, str) and re.search(r'F\s*$', v, re.I):
            return round((numeric(v)-32)*5/9, 8)
        return numeric(v)
    if field in DATE_FIELDS or field.endswith('_utc'):
        if isinstance(v, str):
            v = datetime.fromisoformat(v.replace('Z', '+00:00'))
        if isinstance(v, datetime):
            return v.date().isoformat() if field in DATE_FIELDS else v.isoformat(timespec='microseconds').rstrip('0').rstrip('.')+'Z'
        if isinstance(v, date):
            return v.isoformat()
    return v

wb = load_workbook(SOURCE, read_only=True, data_only=True)
for ws in wb.worksheets:
    rows = ws.iter_rows()
    header = [c.value for c in next(rows)]
    assert all(isinstance(h,str) for h in header)
    records = []
    by_key = {}
    duplicate_count = 0
    changed_count = Counter()
    for cells in rows:
        if all(c.value is None for c in cells):
            continue
        rowno = cells[0].row
        raw = {field: plain(cell.value) for field, cell in zip(header, cells)}
        record = {field: normalize(ws.title,field,cell.value) for field,cell in zip(header,cells)}
        changes = {field: {'from':raw[field], 'to':record[field]} for field in header if raw[field] != record[field]}
        # ISO formatting of existing typed dates is serialization, not a source repair.
        substantive = {k:v for k,v in changes.items() if k not in DATE_FIELDS and not k.endswith('_utc')}
        changed_count.update(substantive.keys())
        key = tuple(record[k] for k in KEYS[ws.title]) if ws.title in KEYS else (rowno,)
        if key in by_key:
            existing = by_key[key]
            assert all(existing[f] == record[f] for f in header), f'Conflicting duplicate in {ws.title}: {key}'
            existing['source']['duplicateRows'].append(rowno)
            duplicate_count += 1
            continue
        record['source'] = {'sheet':ws.title,'row':rowno,'duplicateRows':[],
                            'cells':{f:get_column_letter(i+1)+str(rowno) for i,f in enumerate(header)},
                            'rawValues':raw}
        if substantive:
            record['source']['normalizations'] = substantive
        records.append(record)
        by_key[key] = record
    tables[ws.title] = records
    schema.append({'sheet':ws.title,'rawRows':len(records)+duplicate_count,'cleanRows':len(records),
                   'duplicatesRemoved':duplicate_count,'keyFields':KEYS.get(ws.title,[]),'fields':header,
                   'columnLetters':{f:get_column_letter(i+1) for i,f in enumerate(header)}})
    cleaning.append({'sheet':ws.title,'verifiedDuplicateRowsRemoved':duplicate_count,
                     'normalizedFieldCountsBeforeDeduplication':dict(changed_count)})
wb.close()

assets = tables['Assets']
asset_by_id = {a['asset_id']:a for a in assets}
production = tables['Production']
events = tables['Events']
hse = tables['HSE']
readings = tables['Readings']
options = tables['Solutions']
readme = tables['ReadMe']
dictionary = tables['Dictionary']
readme_by_topic = {r['Topic']:r for r in readme}

def cell(record, field):
    return record['source']['sheet']+'!'+record['source']['cells'][field]

def derived(record, field, value, formula, refs):
    record[field] = value
    record.setdefault('derivedFields',{})[field] = {'formula':formula,'sourceCells':refs}

def enrich_site(record):
    a = asset_by_id[record['asset_id']]
    if 'site_id' not in record:
        derived(record,'site_id',a['site_id'],'Assets lookup on normalized asset_id',[cell(a,'asset_id'),cell(a,'site_id')])
    assert record['site_id'] == a['site_id']
    derived(record,'site_name',a['site_name'],'Assets lookup on normalized asset_id',[cell(a,'asset_id'),cell(a,'site_name')])
    return a

for p in production:
    a = enrich_site(p)
    if p['potential_oil_bbl'] is None:
        derived(p,'potential_oil_bbl',p['calendar_hours']*a['modeled_capacity_bbl_hr'],
                'calendar_hours × modeled_capacity_bbl_hr',[cell(p,'calendar_hours'),cell(a,'modeled_capacity_bbl_hr')])
    derived(p,'net_deferred_bbl',p['gross_deferred_bbl']-p['recovered_oil_bbl'],
            'gross_deferred_bbl − recovered_oil_bbl',[cell(p,'gross_deferred_bbl'),cell(p,'recovered_oil_bbl')])
prod_lookup = {(p['asset_id'],p['month']):p for p in production}

for e in events:
    a = enrich_site(e)
    month = e['start_utc'][:7]+'-01'
    derived(e,'month',month,'Month of start_utc',[cell(e,'start_utc')])
    if e['stopped_hours'] is None:
        assert all(e[k] is not None for k in STAGES)
        derived(e,'stopped_hours',sum(e[k] for k in STAGES),'Sum of five downtime stages',[cell(e,k) for k in STAGES])
    if e['lost_oil_bbl'] is None:
        derived(e,'lost_oil_bbl',e['stopped_hours']*a['modeled_capacity_bbl_hr'],
                'stopped_hours × modeled_capacity_bbl_hr',[cell(e,'stopped_hours'),cell(a,'modeled_capacity_bbl_hr')])
    p = prod_lookup[(e['asset_id'],month)]
    derived(e,'net_lost_bbl',e['lost_oil_bbl']*(1-e['recoverable_fraction']),
            'lost_oil_bbl × (1 − recoverable_fraction)',[cell(e,'lost_oil_bbl'),cell(e,'recoverable_fraction')])
    derived(e,'modeled_contribution_loss_usd',e['net_lost_bbl']*(p['price_usd_bbl']-p['variable_cost_usd_bbl']),
            'lost_oil_bbl × (1 − recoverable_fraction) × (matching asset-month price − variable cost)',
            [cell(e,'lost_oil_bbl'),cell(e,'recoverable_fraction'),cell(p,'price_usd_bbl'),cell(p,'variable_cost_usd_bbl')])
    known_components = e['labor_hours']*e['labor_rate_usd_hr']+e['contractor_usd']
    if e['parts_usd'] is not None:
        known_components += e['parts_usd']
    derived(e,'known_maintenance_cost_usd',known_components,
            'labor_hours × labor_rate_usd_hr + contractor_usd + known parts_usd; pending part cost excluded and unknown',
            [cell(e,k) for k in ['labor_hours','labor_rate_usd_hr','contractor_usd','parts_usd']])
    derived(e,'maintenance_cost_complete',e['parts_usd'] is not None,
            'parts_usd is present; missing invoices remain unknown',[cell(e,'parts_usd'),cell(e,'cost_status')])

for h in hse:
    enrich_site(h)
    derived(h,'is_overdue',h['action_status']=='Open' and h['due_date'] < h['status_as_of'],
            'action_status = Open and due_date < status_as_of',[cell(h,k) for k in ['action_status','due_date','status_as_of']])
    derived(h,'blocks_pilot',h['action_status']=='Open' and h['potential_severity']==5,
            'Workbook HSE gate: open potential_severity 5 blocks affected site pilot',
            [cell(h,'action_status'),cell(h,'potential_severity'),cell(readme_by_topic['HSE decision gate'],'Instructions')])
    derived(h,'unrecovered_spill_bbl',h['spilled_oil_bbl']-h['recovered_spill_bbl'],
            'spilled_oil_bbl − recovered_spill_bbl',[cell(h,'spilled_oil_bbl'),cell(h,'recovered_spill_bbl')])

for r in readings:
    enrich_site(r)
    lag = (datetime.fromisoformat(r['received_utc'].replace('Z','+00:00'))-datetime.fromisoformat(r['observed_utc'].replace('Z','+00:00'))).total_seconds()/3600
    assert lag >= 0
    derived(r,'ingestion_delay_hours',lag,'(received_utc − observed_utc) in hours',[cell(r,'received_utc'),cell(r,'observed_utc')])

def total(records, field):
    values = [r[field] for r in records]
    assert all(v is not None for v in values), field
    return float(sum(Decimal(str(v)) for v in values))

def event_totals(records, breakdown=True):
    d = {'event_count':len(records),'stopped_hours':total(records,'stopped_hours'),
         'stages':{k:total(records,k) for k in STAGES},
         'parts_wait_event_count':sum(r['parts_wait_hours']>0 for r in records),
         'gross_deferred_bbl':total(records,'lost_oil_bbl'),'net_lost_bbl':total(records,'net_lost_bbl'),
         'modeled_contribution_loss_usd':total(records,'modeled_contribution_loss_usd'),
         'known_maintenance_cost_usd':total(records,'known_maintenance_cost_usd'),
         'pending_parts_invoice_count':sum(not r['maintenance_cost_complete'] for r in records),
         'source':{'sheet':'Events','rows':[r['source']['row'] for r in records],'ids':[r['event_id'] for r in records]}}
    if breakdown:
        d['cause_breakdown'] = [dict(cause=c,**event_totals([r for r in records if r['cause']==c],False)) for c in sorted({r['cause'] for r in records})]
        d['months'] = [dict(month=m,**event_totals([r for r in records if r['month']==m],False)) for m in sorted({r['month'] for r in records})]
    return d

def production_totals(records, breakdown=True):
    d = {'record_count':len(records),**{k:total(records,k) for k in ['potential_oil_bbl','produced_oil_bbl','gross_deferred_bbl','recovered_oil_bbl','net_deferred_bbl']},
         'source':{'sheet':'Production','rows':[r['source']['row'] for r in records]}}
    d['produced_share_of_potential_pct'] = d['produced_oil_bbl']/d['potential_oil_bbl']*100 if d['potential_oil_bbl'] else None
    if breakdown:
        d['months'] = [dict(month=m,**production_totals([r for r in records if r['month']==m],False)) for m in sorted({r['month'] for r in records})]
    return d

def hse_totals(records):
    opened = [r for r in records if r['action_status']=='Open']
    return {'observation_count':len(records),'open_actions':len(opened),'closed_actions':sum(r['action_status']=='Closed' for r in records),
            'open_critical_actions':sum(r['blocks_pilot'] for r in records),'overdue_open_actions':sum(r['is_overdue'] for r in records),
            'open_action_setup_usd':total(opened,'action_setup_usd'),'open_action_annual_usd':total(opened,'action_annual_usd'),
            'open_action_person_hours':total(opened,'action_person_hours'),'spilled_oil_bbl':total(records,'spilled_oil_bbl'),
            'recovered_spill_bbl':total(records,'recovered_spill_bbl'),'unrecovered_spill_bbl':total(records,'unrecovered_spill_bbl'),
            'hazard_counts':dict(Counter(r['hazard'] for r in records)),
            'source':{'sheet':'HSE','rows':[r['source']['row'] for r in records],'ids':[r['hse_id'] for r in records]}}

def reading_totals(records):
    lags = [r['ingestion_delay_hours'] for r in records]
    return {'reading_count':len(records),'quality_counts':dict(Counter(r['quality'] for r in records)),
            'missing_vibration_count':sum(r['vibration_mm_s'] is None for r in records),
            'missing_temperature_count':sum(r['temperature_c'] is None for r in records),
            'delayed_over_1h_count':sum(t>1 for t in lags),
            'mean_ingestion_delay_hours':sum(lags)/len(lags) if lags else None,
            'max_ingestion_delay_hours':max(lags) if lags else None,
            'source':{'sheet':'Readings','rows':[r['source']['row'] for r in records]}}

def summaries(site_id=None, asset_id=None):
    def selected(rows):
        return [r for r in rows if (site_id is None or r['site_id']==site_id) and (asset_id is None or r['asset_id']==asset_id)]
    selected_events = selected(events)
    return {'unplanned':event_totals([r for r in selected_events if r['event_type']=='Unplanned']),
            'planned':event_totals([r for r in selected_events if r['event_type']=='Planned']),
            'production':production_totals(selected(production)),'hse':hse_totals(selected(hse)),
            'readings':reading_totals(selected(readings))}

SITE_FIELDS = ['network_uptime_pct','spare_fill_pct','technicians','training_hours_per_person','available_project_hours','cleanup_setup_usd','edge_setup_usd','edge_annual_usd']
sites = []
for site_id in sorted({a['site_id'] for a in assets}):
    site_assets = [a for a in assets if a['site_id']==site_id]
    a = site_assets[0]
    assert all(all(s[f]==a[f] for f in SITE_FIELDS+['site_name']) for s in site_assets)
    s = {'site_id':site_id,'site_name':a['site_name'],'asset_ids':[x['asset_id'] for x in site_assets],
         **{f:a[f] for f in SITE_FIELDS},'training_required_hours':a['technicians']*a['training_hours_per_person'],
         **summaries(site_id=site_id),
         'source':{'sheet':'Assets','row':a['source']['row'],'repeatedSiteRows':[x['source']['row'] for x in site_assets],
                   'cells':{f:a['source']['cells'][f] for f in ['site_id','site_name']+SITE_FIELDS}},
         'derivedFields':{'training_required_hours':{'formula':'technicians × training_hours_per_person',
                  'sourceCells':[cell(a,'technicians'),cell(a,'training_hours_per_person'),cell(readme_by_topic['Implementation and training'],'Instructions')]}}}
    sites.append(s)
for a in assets:
    a['metrics'] = summaries(asset_id=a['asset_id'])
company = summaries()
company['site_count'] = len(sites)
company['asset_count'] = len(assets)
company['workforce'] = {'technicians':sum(s['technicians'] for s in sites),
                        'training_required_hours':sum(s['training_required_hours'] for s in sites),
                        'available_project_hours':sum(s['available_project_hours'] for s in sites),
                        'source':{'sheet':'Assets','rows':[s['source']['row'] for s in sites],
                                  'rule':'Site fields counted once per site, not once per asset'}}

# Independent reconciliation against the monthly production control table.
reconciliation = {'max_event_stage_difference_hours':0,'max_event_capacity_difference_bbl':0,
                  'max_monthly_gross_difference_bbl':0,'max_monthly_net_difference_bbl':0,
                  'max_production_identity_difference_bbl':0}
for e in events:
    reconciliation['max_event_stage_difference_hours'] = max(reconciliation['max_event_stage_difference_hours'],abs(e['stopped_hours']-sum(e[k] for k in STAGES)))
    reconciliation['max_event_capacity_difference_bbl'] = max(reconciliation['max_event_capacity_difference_bbl'],abs(e['lost_oil_bbl']-e['stopped_hours']*asset_by_id[e['asset_id']]['modeled_capacity_bbl_hr']))
for p in production:
    linked = [e for e in events if (e['asset_id'],e['month'])==(p['asset_id'],p['month'])]
    reconciliation['max_monthly_gross_difference_bbl'] = max(reconciliation['max_monthly_gross_difference_bbl'],abs(total(linked,'lost_oil_bbl')-p['gross_deferred_bbl']))
    reconciliation['max_monthly_net_difference_bbl'] = max(reconciliation['max_monthly_net_difference_bbl'],abs(total(linked,'net_lost_bbl')-p['net_deferred_bbl']))
    reconciliation['max_production_identity_difference_bbl'] = max(reconciliation['max_production_identity_difference_bbl'],abs(p['potential_oil_bbl']-p['gross_deferred_bbl']+p['recovered_oil_bbl']-p['produced_oil_bbl']))
assert reconciliation['max_event_stage_difference_hours'] < 1e-8
assert reconciliation['max_event_capacity_difference_bbl'] < 1e-8
assert reconciliation['max_monthly_gross_difference_bbl'] < 1e-8
assert reconciliation['max_monthly_net_difference_bbl'] <= 0.00500001
assert reconciliation['max_production_identity_difference_bbl'] < 0.011
assert len(sites)==6 and len(assets)==18 and len(events)==385 and len(production)==216 and len(hse)==36 and len(readings)==6570

limitations = [
    'All records are from the fictional case workbook; historical case values are not live operational data.',
    'HSE status is the 2026-10-01 source snapshot. Open severity-5 actions block the affected site pilot; absence of a critical hold is not an authorization record.',
    'No inventory parts register, stock quantity, reorder point, supplier, purchase request or receipt records exist in the workbook.',
    'Training fields specify required effort and solution-option costs. No employee identities, attendance, completion, configuration-task or course records are provided.',
    'No user accounts, permissions, audit-login history, security-incident records or cyber-risk measurements are provided.',
    'Solutions are workbook scenario options, not approved purchases, deployments or observed effectiveness. No intervention forecast is calculated in this export.',
    'Network uptime is connectivity, not production availability. Missing readings remain null; Suspect is not a confirmed equipment fault.',
    'Known maintenance costs omit 14 unknown parts invoices. Do not present the subtotal as a complete total.',
    'Modeled contribution loss is calculated from the workbook method, not a source financial transaction. Monthly production loss must not be added to event losses.',
    'HSE-event linkage is context, not evidence that the hazard caused the outage. Cause categories overlap downtime stages.',
    'Severity and likelihood are ordinal case bands. No injury-rate denominator or proven incident-prevention benefit is supplied.',
]
data = {'meta':{'source':SOURCE.name,'source_path':str(SOURCE),'source_sha256':sha256(SOURCE.read_bytes()).hexdigest(),
                'period':'2025','hse_status_as_of':sorted({h['status_as_of'] for h in hse}),
                'cleaning':cleaning,'limitations':limitations,'reconciliation':reconciliation,
                'dataAvailability':{'inventory':{'present':['Assets.spare_fill_pct','Events.parts_wait_hours','Events.parts_usd'],
                    'absent':['part IDs','part names','stock quantities','reorder thresholds','supplier lead times','purchasing requests','stock receipts']},
                    'training':{'present':['Assets.technicians','Assets.training_hours_per_person','Assets.available_project_hours','Solutions O6 Digital training'],
                    'absent':['employee names','training assignments','attendance','completion status','configuration progress']}}},
        'sheetSchema':schema,'company':company,'sites':sites,'assets':assets,'events':events,'hse':hse,
        'production':production,'readings':readings,'options':options,'readme':readme,'dictionary':dictionary}
OUT.joinpath('workbook-data.json').write_text(json.dumps(data,separators=(',',':'),ensure_ascii=False,allow_nan=False)+'\n')

lines = ['# Workbook-only data audit','',f'Source: {SOURCE.name}',f'SHA-256: {data["meta"]["source_sha256"]}',
         '','All exported business records come solely from the supplied workbook. The source workbook was not changed.',
         '','## Source tables','', '| Sheet | Raw rows | Unique rows | Removed verified copies |', '|---|---:|---:|---:|']
for s in schema:
    lines.append(f'| {s["sheet"]} | {s["rawRows"]:,} | {s["cleanRows"]:,} | {s["duplicatesRemoved"]:,} |')
lines += ['', '## Company totals', '', f'- Unplanned downtime: {company["unplanned"]["stopped_hours"]:,.2f} equipment-hours across {company["unplanned"]["event_count"]} events.',
          f'- Planned downtime: {company["planned"]["stopped_hours"]:,.2f} equipment-hours across {company["planned"]["event_count"]} events.',
          f'- HSE: {company["hse"]["open_actions"]} open actions, {company["hse"]["open_critical_actions"]} open severity-5 actions, {company["hse"]["overdue_open_actions"]} overdue as of 2026-10-01.',
          f'- Workforce: {company["workforce"]["technicians"]} technicians, {company["workforce"]["training_required_hours"]} required training hours, {company["workforce"]["available_project_hours"]} available project hours.',
          '', '## Site totals', '', '| Site | Unplanned h | Parts wait h | Open HSE | Critical HSE | Fill % | Network % | Technicians | Training h | Project h |', '|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|']
for s in sites:
    lines.append(f'| {s["site_name"]} | {s["unplanned"]["stopped_hours"]:,.2f} | {s["unplanned"]["stages"]["parts_wait_hours"]:,.2f} | {s["hse"]["open_actions"]} | {s["hse"]["open_critical_actions"]} | {s["spare_fill_pct"]} | {s["network_uptime_pct"]} | {s["technicians"]} | {s["training_required_hours"]} | {s["available_project_hours"]} |')
lines += ['', '## Provenance and transformations', '',
          'Each business record includes source.sheet, source.row, source.cells, source.rawValues and source.duplicateRows. Normalized values retain their original values and transformation. Derived fields state the formula and source cells. Aggregate source objects list contributing source rows.',
          'IDs are trimmed, uppercased and padded where needed; numeric unit strings are parsed; explicit Fahrenheit readings are converted to Celsius. Four missing event durations and three missing production potentials are restored only by the workbook formulas. Original nulls remain in source.rawValues. All missing sensor measurements and pending parts invoices remain null.',
          'Dates retain source meaning: event and reading timestamps are UTC; due dates and HSE snapshot dates are calendar dates. No current-date overdue calculation is used.',
          'Duplicate keys were checked for field equality after normalization. No conflicting duplicate was found.',
          '', '## Reconciliation', '', '```json', json.dumps(reconciliation,indent=2), '```',
          'Gross event losses reconcile to monthly gross deferred oil. Net event losses differ by at most 0.005 bbl because monthly recovered oil is rounded. Event stage sums and capacity calculations reconcile.',
          '', '## Remove from the previous interface', '',
          '- Synthetic part names, DM-* part IDs, stock balances, reorder points, reservations, purchasing requests and receipts.',
          '- Synthetic employee records, training completion claims, configuration-task progress and fictional activity logs.',
          '- Simulated HSE closures or pilot authorizations presented as source statuses. Source HSE records must stay unchanged.',
          '- The $8,000 configuration / 80 external-hour allowance, $1,500 support allowance, combined $58,500/$6,200 budget, and 100-hour selected plan. These are prior proposal choices, not approved workbook records.',
          '- Scenario sliders, forecast downtime reductions and forecast benefits from the operational-data views. Workbook option coefficients may remain as labeled option reference only.',
          '- Assertions that SSO, production authorization, protected audit logging, equipment control or offline synchronization exist. A permission design may be shown separately as proposed app design.',
          '', '## Limits', '']
lines += ['- '+t for t in limitations]
lines += ['', '## Sheet fields', '']
for s in schema:
    lines += ['### '+s['sheet'], '', ', '.join(s['fields']), '']
OUT.joinpath('workbook-data-audit.md').write_text('\n'.join(lines)+'\n')
print(json.dumps({'output':str(OUT/'workbook-data.json'),'bytes':(OUT/'workbook-data.json').stat().st_size,
                  'tables':[{k:s[k] for k in ['sheet','rawRows','cleanRows','duplicatesRemoved']} for s in schema],
                  'unplanned_hours':company['unplanned']['stopped_hours'],'hse':company['hse'],
                  'workforce':company['workforce'],'checks':reconciliation},indent=2))
